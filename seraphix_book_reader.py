"""Legacy read-only DOCX extraction and chunk indexing.

Usage: python seraphix_book_reader.py ./examples/synthetic_manuscript.docx
The example path is illustrative; no manuscript is distributed.
Outputs go to Memory/Book_Index/chunks.json, never to the source DOCX.
"""
from __future__ import annotations
import argparse
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from docx import Document


def extract(docx_path: Path) -> list[dict]:
    doc = Document(str(docx_path))
    blocks = []
    for index, paragraph in enumerate(doc.paragraphs, 1):
        value = paragraph.text.strip()
        if value:
            blocks.append({"paragraph": index, "style": paragraph.style.name if paragraph.style else "", "text": value})
    for table_number, table in enumerate(doc.tables, 1):
        for row_number, row in enumerate(table.rows, 1):
            value = " | ".join(cell.text.strip().replace("\n", " / ") for cell in row.cells)
            if value.strip(" |"):
                blocks.append({"table": table_number, "row": row_number, "text": value})
    return blocks


def chunk(blocks: list[dict], max_chars: int = 2800) -> list[dict]:
    chunks = []
    buffer = []
    count = 0
    start = None
    for block in blocks:
        text = block["text"]
        # Split exceptionally long paragraphs safely, without dropping text.
        pieces = [text[i:i + max_chars] for i in range(0, len(text), max_chars)]
        for piece in pieces:
            if buffer and count + len(piece) + 2 > max_chars:
                chunks.append({"id": len(chunks) + 1, "start": start, "text": "\n\n".join(buffer)})
                buffer, count, start = [], 0, None
            if start is None:
                start = {k: v for k, v in block.items() if k != "text"}
            buffer.append(piece)
            count += len(piece) + 2
    if buffer:
        chunks.append({"id": len(chunks) + 1, "start": start, "text": "\n\n".join(buffer)})
    return chunks


def main():
    parser = argparse.ArgumentParser(description="Read-only DOCX indexer for Seraphix")
    parser.add_argument("docx", type=Path)
    args = parser.parse_args()
    source = args.docx.expanduser().resolve()
    if not source.is_file() or source.suffix.lower() != ".docx":
        parser.error("Provide an existing .docx file")
    blocks = extract(source)
    chunks = chunk(blocks)
    output = Path(__file__).resolve().parent / "Memory" / "Book_Index"
    output.mkdir(parents=True, exist_ok=True)
    payload = {
        "source": str(source),
        "indexed_at_utc": datetime.now(timezone.utc).isoformat(),
        "paragraphs_and_table_rows": len(blocks),
        "chunks": chunks,
    }
    destination = output / "chunks.json"
    destination.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"Indexed {len(blocks)} text blocks into {len(chunks)} chunks")
    print(f"Saved: {destination}")
    print("Source manuscript was not modified. This indexes text; it does not yet add book Q&A to the GUI.")

if __name__ == "__main__":
    main()
